from docx import Document
from report_export import build_report_docx


def _report(**overrides):
    base = {
        "title": "Test Report",
        "reporting_period": "Q1 2026",
        "executive_summary": "Summary text [1].",
        "sections": [
            {"heading": "1. Findings", "body": "Body paragraph.", "key_stat": "42%"},
        ],
        "risk_matrix": [],
        "conclusions": "Conclusion text.",
        "recommendations": [],
        "data_coverage": {},
        "sources": [],
    }
    base.update(overrides)
    return base


def _read_back(docx_bytes: bytes) -> Document:
    import io
    return Document(io.BytesIO(docx_bytes))


def test_build_report_docx_returns_nonempty_bytes():
    docx_bytes = build_report_docx(_report(), "What happened?")
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 0


def test_uses_report_title_over_the_raw_query():
    doc = _read_back(build_report_docx(_report(title="A Formal Title"), "raw user query"))
    assert doc.paragraphs[0].text == "A Formal Title"
    assert doc.paragraphs[0].style.name == "Title"


def test_falls_back_to_query_when_report_has_no_title():
    doc = _read_back(build_report_docx(_report(title=None), "raw user query"))
    assert doc.paragraphs[0].text == "raw user query"


def test_strips_markdown_bold_and_code_spans():
    report = _report(executive_summary="A **bold** claim about `metric_name`.")
    doc = _read_back(build_report_docx(report, "q"))
    texts = [p.text for p in doc.paragraphs]
    assert "A bold claim about metric_name." in texts
    # The markdown punctuation itself must not leak into the visible text.
    assert not any("**" in t or "`" in t for t in texts)


def test_bold_span_becomes_a_real_bold_run():
    report = _report(executive_summary="Plain **bold** plain.")
    doc = _read_back(build_report_docx(report, "q"))
    summary_para = next(p for p in doc.paragraphs if "bold" in p.text)
    bold_runs = [r for r in summary_para.runs if r.bold]
    assert any(r.text == "bold" for r in bold_runs)


def test_multi_paragraph_section_body_becomes_separate_paragraphs():
    report = _report(sections=[
        {"heading": "H", "body": "First para.\n\nSecond para.", "key_stat": ""},
    ])
    doc = _read_back(build_report_docx(report, "q"))
    texts = [p.text for p in doc.paragraphs]
    assert "First para." in texts
    assert "Second para." in texts


def test_risk_matrix_becomes_a_real_table_with_dynamic_dimension_columns():
    report = _report(risk_matrix=[
        {"entity": "Bank A", "dimensions": {"Risk": "High", "Trend": "Low"},
         "overall": "High", "priority_action": "Review"},
    ])
    doc = _read_back(build_report_docx(report, "q"))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Entity", "Risk", "Trend", "Overall", "Priority Action"]
    assert [c.text for c in table.rows[1].cells] == ["Bank A", "High", "Low", "High", "Review"]


def test_no_risk_matrix_section_when_empty():
    doc = _read_back(build_report_docx(_report(risk_matrix=[]), "q"))
    assert len(doc.tables) == 0
    assert not any("Risk Matrix" == p.text for p in doc.paragraphs)


def test_recommendations_become_bulleted_list_items():
    report = _report(recommendations=["Do this.", "Do that."])
    doc = _read_back(build_report_docx(report, "q"))
    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in bullets] == ["Do this.", "Do that."]


def test_sources_are_listed_with_their_citation_id():
    report = _report(sources=[{"id": "1", "question": "Q1?", "summary": "Answer summary."}])
    doc = _read_back(build_report_docx(report, "q"))
    texts = [p.text for p in doc.paragraphs]
    assert "[1] Q1?" in texts
    assert "Answer summary." in texts


def test_data_coverage_line_includes_formatted_record_count():
    report = _report(data_coverage={
        "sub_questions_answered": 3, "total_records_analysed": 250000,
        "datasets_used": ["transactions_data.csv"],
    })
    doc = _read_back(build_report_docx(report, "q"))
    texts = [p.text for p in doc.paragraphs]
    assert any("250,000 records analysed" in t for t in texts)
    assert any("3 sub-questions answered" in t for t in texts)
    assert any("transactions_data.csv" in t for t in texts)


def test_classification_is_rendered_uppercased():
    doc = _read_back(build_report_docx(_report(classification="Confidential — Internal Use"), "q"))
    texts = [p.text for p in doc.paragraphs]
    assert "CONFIDENTIAL — INTERNAL USE" in texts
